# -*- coding: utf-8 -*-
"""
data/자원봉사.xlsx -> data/volunteer.json   (탭데이터형)
시트 '자원봉사', 머리글 행(기본 4행): 게시 / 탭 / 종류 / 이름·라벨 / 내용 / 태그 / 아이콘 / 사진 / 캡션
- 탭 빈칸 = 상단: 소개(문구) · 신청(이름=라벨, 내용=번호/안내, 아이콘)
- 탭 이름 있음 = 그 탭 안: 카드 · 절차 · 혜택 · 안내 · 사진(사진=파일명, 캡션)
오류가 있으면 JSON을 만들지 않고 멈춘다(기존 화면 유지).
"""
import json, sys
from datetime import datetime, timezone, timedelta
from pathlib import Path
import openpyxl

SRC = Path("data/자원봉사.xlsx"); OUT = Path("data/volunteer.json")
KST = timezone(timedelta(hours=9)); SHEET = "자원봉사"
TOP_KINDS = {"소개","신청"}
TAB_KINDS = {"카드","절차","혜택","안내","사진"}
ICONS = {
 "phone","home","meal","talk","user","users","clean","music","health","pill","walk","shield","star",
 "heart","hand","gift","bread","coffee","mic","mega","bell","mail","book","edit","camera","calendar",
 "clock","car","bus","mappin","leaf","flower","sun","smile","scissors","wrench","bank","card","box",
 "link","receipt","pin","route","check"}

def s(v): return "" if v is None else str(v).strip()

def main():
    if not SRC.exists(): print(f"::error::{SRC} 파일이 없습니다."); sys.exit(1)
    wb = openpyxl.load_workbook(SRC, data_only=True)
    if SHEET not in wb.sheetnames: print(f"::error::'{SHEET}' 시트가 없습니다."); sys.exit(1)
    ws = wb[SHEET]
    head_row = 4
    for r in range(1, min(ws.max_row, 10)+1):
        vals = [s(c.value) for c in ws[r]]
        if "종류" in vals and "게시" in vals: head_row = r; break
    head = [s(c.value) for c in ws[head_row]]
    errors = [f"{head_row}행에 '{h}' 칸이 없습니다." for h in ["게시","탭","종류","이름·라벨","내용"] if h not in head]
    if errors:
        for e in errors: print("::error::"+e)
        sys.exit(1)
    ix = {h:i for i,h in enumerate(head)}
    def col(row, n):
        i = ix.get(n); return s(row[i]) if (i is not None and i < len(row)) else ""

    doc = {"생성시각": datetime.now(KST).isoformat(timespec="seconds"), "출처": SRC.name,
           "상단": {"소개":"", "신청":[]}, "탭": []}
    tabs, order = {}, []
    for r in range(head_row+1, ws.max_row+1):
        row = [c.value for c in ws[r]]
        if not any(s(v) for v in row): continue
        if col(row,"게시").upper() in ("X","×"): continue
        탭, 종류 = col(row,"탭"), col(row,"종류")
        이름, 내용, 태그, 아이콘 = col(row,"이름·라벨"), col(row,"내용"), col(row,"태그"), col(row,"아이콘")
        사진, 캡션 = col(row,"사진"), col(row,"캡션")
        if not 종류: errors.append(f"{r}행: 종류가 비었습니다."); continue
        if 아이콘 and 아이콘 not in ICONS: errors.append(f"{r}행: 아이콘 '{아이콘}'은 목록에 없습니다."); continue
        if not 탭:
            if 종류 not in TOP_KINDS:
                errors.append(f"{r}행: '{종류}'는 탭 안에 들어가는 항목이라 탭 이름이 필요합니다."); continue
            if 종류 == "소개": doc["상단"]["소개"] = doc["상단"]["소개"] or 내용
            else: doc["상단"]["신청"].append({"라벨":이름, "값":내용, "아이콘":아이콘})
            continue
        if 종류 not in TAB_KINDS:
            errors.append(f"{r}행: '{종류}'는 상단 항목이라 탭 칸을 비워야 합니다."); continue
        if 탭 not in tabs:
            tabs[탭] = {"이름":탭, "카드":[], "절차":[], "혜택":[], "안내":"", "사진":[]}; order.append(탭)
        t = tabs[탭]
        if 종류 == "카드":   t["카드"].append({"이름":이름, "설명":내용, "태그":태그, "아이콘":아이콘})
        elif 종류 == "절차": t["절차"].append(이름 or 내용)
        elif 종류 == "혜택": t["혜택"].append(내용 or 이름)
        elif 종류 == "안내": t["안내"] = t["안내"] or 내용
        elif 종류 == "사진":
            fn = 사진 or 내용
            if not fn: errors.append(f"{r}행: 사진은 파일명이 필요합니다."); continue
            t["사진"].append({"파일":fn, "캡션":캡션})
    if errors:
        for e in errors: print("  ✗ "+e)
        print(f"::error::잘못된 값 {len(errors)}군데 — 기존 JSON 유지"); sys.exit(1)
    doc["탭"] = [tabs[n] for n in order]
    OUT.write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"✓ {OUT} — 탭 {len(doc['탭'])}개, 카드 {sum(len(t['카드']) for t in doc['탭'])}, 사진 {sum(len(t['사진']) for t in doc['탭'])}")

if __name__ == "__main__":
    main()
