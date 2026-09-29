# -*- coding: utf-8 -*-
"""
data/자주묻는질문.xlsx -> data/faq.json
시트 'FAQ', 4행 머리글(게시/분류/질문/답변), 5행부터.
분류: 이용·회원 / 프로그램 / 식사 / 후원 / 자원봉사 / 기타
답변 안에서 줄바꿈(엔터)은 화면에서도 줄바꿈으로 표시됩니다.
"""
import json, sys
from datetime import datetime, timezone, timedelta
from pathlib import Path
import openpyxl

SRC=Path("data/자주묻는질문.xlsx"); OUT=Path("data/faq.json")
KST=timezone(timedelta(hours=9)); SHEET="FAQ"
ORDER=["이용·회원","프로그램","식사","후원","자원봉사","기타"]
def s(v): return "" if v is None else str(v).strip()

def main():
    if not SRC.exists(): print(f"::error::{SRC} 없음"); sys.exit(1)
    wb=openpyxl.load_workbook(SRC,data_only=True)
    if SHEET not in wb.sheetnames: print(f"::error::'{SHEET}' 시트 없음"); sys.exit(1)
    ws=wb[SHEET]
    head=[s(c.value) for c in ws[4]]
    need=["게시","분류","질문","답변"]; errors=[]
    for h in need:
        if h not in head: errors.append(f"4행에 '{h}' 칸 없음")
    if errors:
        for e in errors: print("::error::"+e)
        sys.exit(1)
    ix={h:head.index(h) for h in head}
    def col(row,n):
        i=ix.get(n); return s(row[i]) if (i is not None and i<len(row)) else ""
    items=[]; cats=[]
    for r in range(5,ws.max_row+1):
        row=[c.value for c in ws[r]]
        if not any(s(v) for v in row): continue
        if col(row,"게시").upper() in ("X","×"): continue
        cat=col(row,"분류"); q=col(row,"질문"); a=col(row,"답변")
        if not q: continue
        if cat and cat not in cats: cats.append(cat)
        items.append({"cat":cat or "기타","q":q,"a":a})
    if not items:
        print("::error::게시할 질문이 한 건도 없습니다."); sys.exit(1)
    cat_sorted=[c for c in ORDER if c in cats]+[c for c in cats if c not in ORDER]
    doc={"생성시각":datetime.now(KST).isoformat(timespec="seconds"),"출처":SRC.name,
         "categories":cat_sorted,"items":items}
    OUT.write_text(json.dumps(doc,ensure_ascii=False,indent=1),encoding="utf-8")
    print(f"✓ {OUT} — 문항 {len(items)} / 분류 {len(cat_sorted)}")

if __name__=="__main__": main()
