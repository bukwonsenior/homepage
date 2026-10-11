# -*- coding: utf-8 -*-
"""
data/조직현황.xlsx  ->  data/org.json
시트: 직원(게시/팀/이름/직책/인원구분/주요업무/비고) · 인원현황(인원구분) · 조직도(게시/구분/상자이름/소제목/항목) · 운영위원회(게시/이름/소속/비고)
모든 시트 4행 = 머리글, 5행부터 데이터. 인원은 직원 시트에서 직접 센다(엑셀 수식 값에 의존하지 않음).
"""
import json, sys
from datetime import datetime, timezone, timedelta
from pathlib import Path
import openpyxl

SRC = Path("data/조직현황.xlsx")
OUT = Path("data/org.json")
KST = timezone(timedelta(hours=9))
TABS = ["조직도", "직원 및 주요업무", "운영위원회"]
KINDS = ["상단", "관리", "왼쪽", "오른쪽", "팀"]


def s(v):
    return "" if v is None else str(v).strip()


def sheet_rows(wb, name, need, errors):
    if name not in wb.sheetnames:
        errors.append(f"'{name}' 시트가 없습니다. 시트 이름을 바꾸지 마세요.")
        return []
    ws = wb[name]
    head = [s(c.value) for c in ws[4]]
    miss = [h for h in need if h not in head]
    if miss:
        errors.append(f"'{name}' 시트 4행에 {', '.join(miss)} 칸이 없습니다.")
        return []
    ix = {h: head.index(h) for h in need}
    out = []
    for r in range(5, ws.max_row + 1):
        row = [c.value for c in ws[r]]
        d = {h: (s(row[i]) if i < len(row) else "") for h, i in ix.items()}
        if not any(d.values()):
            continue
        if "게시" in d and d["게시"].upper() in ("X", "×"):
            continue
        d["_행"] = r
        out.append(d)
    return out


def main():
    if not SRC.exists():
        print(f"::error::{SRC} 파일이 없습니다."); sys.exit(1)
    wb = openpyxl.load_workbook(SRC, data_only=True)
    errors, warns = [], []

    staff = sheet_rows(wb, "직원", ["게시", "팀", "이름", "직책", "인원구분", "주요업무"], errors)
    kinds = sheet_rows(wb, "인원현황", ["인원구분"], errors)
    org = sheet_rows(wb, "조직도", ["게시", "구분", "상자이름", "소제목", "항목"], errors)
    comm = sheet_rows(wb, "운영위원회", ["게시", "이름", "소속", "비고"], errors)
    if errors:
        for e in errors: print("::error::" + e)
        sys.exit(1)

    # 직원 → 팀별 묶음 (나온 순서)
    teams, tmap = [], {}
    for d in staff:
        if not d["팀"]: errors.append(f"직원 {d['_행']}행: 팀 칸이 비었습니다"); continue
        if not d["이름"]: errors.append(f"직원 {d['_행']}행: 이름 칸이 비었습니다"); continue
        if d["팀"] not in tmap:
            tmap[d["팀"]] = {"이름": d["팀"], "사람": []}
            teams.append(tmap[d["팀"]])
        tmap[d["팀"]]["사람"].append({"이름": d["이름"], "직책": d["직책"], "업무": d["주요업무"]})

    # 인원표 (인원현황 순서, 목록에 없는 구분은 뒤에 붙임)
    order = [d["인원구분"] for d in kinds if d["인원구분"] and d["인원구분"] not in ("계", "게시 직원 수", "점검")]
    cnt = {}
    for d in staff:
        if not d["이름"]: continue
        k = d["인원구분"] or d["직책"]
        if not d["인원구분"]: warns.append(f"직원 {d['_행']}행({d['이름']}): 인원구분이 비어 직책({d['직책']})으로 셉니다")
        cnt[k] = cnt.get(k, 0) + 1
        if k not in order:
            order.append(k)
            if d["인원구분"]: warns.append(f"직원 {d['_행']}행: 인원구분 '{k}'이(가) 인원현황 목록에 없어 맨 뒤에 붙였습니다")
    counts = [{"구분": k, "수": cnt.get(k, 0)} for k in order if cnt.get(k, 0) > 0]

    # 조직도
    chart = {"상단": None, "관리": [], "왼쪽": [], "오른쪽": [], "팀": []}
    cmap = {}
    for d in org:
        g, n = d["구분"], d["상자이름"]
        if g not in KINDS: errors.append(f"조직도 {d['_행']}행: 구분은 {'/'.join(KINDS)} 중 하나여야 합니다 (지금: {g})"); continue
        if not n: errors.append(f"조직도 {d['_행']}행: 상자이름이 비었습니다"); continue
        if g == "상단":
            chart["상단"] = {"이름": n, "윗글씨": d["항목"]}
        elif g == "팀":
            if n not in cmap:
                cmap[n] = {"이름": n, "항목": []}
                chart["팀"].append(cmap[n])
            if d["항목"]:
                cmap[n]["항목"].append({"소제목": d["소제목"], "항목": d["항목"]})
        else:
            chart[g].append(n)

    members = []
    for d in comm:
        if not d["이름"]: errors.append(f"운영위원회 {d['_행']}행: 이름 칸이 비었습니다"); continue
        members.append({"이름": d["이름"], "소속": d["소속"], "비고": d["비고"]})

    for w in warns: print("::warning::" + w)
    if errors:
        print("\n고쳐야 할 곳:")
        for e in errors: print("  ✗ " + e)
        print(f"::error::잘못된 값 {len(errors)}군데")
        sys.exit(1)

    doc = {
        "생성시각": datetime.now(KST).isoformat(timespec="seconds"),
        "출처": SRC.name,
        "조직도": chart,
        "인원": counts,
        "계": sum(c["수"] for c in counts),
        "직원": teams,
        "위원": members,
        "탭": [{"이름": t} for t in TABS],
    }
    OUT.write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"✓ {OUT} — 직원 {doc['계']}명(팀 {len(teams)}개), 조직도 팀 {len(chart['팀'])}개, 운영위원 {len(members)}명")


if __name__ == "__main__":
    main()
