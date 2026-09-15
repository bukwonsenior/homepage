# -*- coding: utf-8 -*-
"""
엑셀(북원_홈페이지_프로그램관리.xlsx)  ->  data/programs.json

담당자가 data/ 폴더에 엑셀을 올리면 GitHub Actions가 이 스크립트를 돌린다.
값이 잘못되어 있으면 JSON을 만들지 않고 멈춘다. (잘못된 데이터가 홈페이지에 올라가지 않게)
"""
import json, re, sys, hashlib
from datetime import datetime, date, timezone, timedelta
from pathlib import Path

import openpyxl

XLSX = Path("data/북원_홈페이지_프로그램관리.xlsx")
OUT  = Path("data/programs.json")
KST  = timezone(timedelta(hours=9))

GUBUN = {"평생교육", "동아리", "특강"}
DAYS  = {"월", "화", "수", "목", "금", "토", "일"}
TIME_RE = re.compile(r"^([01]\d|2[0-3]):[0-5]\d$")

# 공개 저장소에 절대 올라가면 안 되는 것
PRIVACY = [
    (re.compile(r"\b01[016-9][-. ]?\d{3,4}[-. ]?\d{4}\b"), "휴대전화번호"),
    (re.compile(r"\b\d{6}[-]\d{7}\b"),                     "주민등록번호"),
]

errors, warns = [], []

def err(row, msg):  errors.append(f"{row}행: {msg}")
def warn(row, msg): warns.append(f"{row}행: {msg}")

def s(v):
    if v is None: return ""
    if isinstance(v, (datetime, date)): return v.strftime("%Y-%m-%d")
    return str(v).strip()

def as_time(v):
    """엑셀이 시간을 datetime.time 으로 바꿔 놓는 경우까지 받아준다."""
    if v is None: return ""
    if hasattr(v, "hour") and hasattr(v, "minute") and not isinstance(v, (datetime, date)):
        return f"{v.hour:02d}:{v.minute:02d}"
    if isinstance(v, datetime): return v.strftime("%H:%M")
    return str(v).strip()

def as_date(v):
    if v is None or s(v) == "": return None
    if isinstance(v, (datetime, date)): return (v.date() if isinstance(v, datetime) else v).isoformat()
    t = s(v).replace(".", "-").replace("/", "-").strip("-")
    try: return date.fromisoformat(t).isoformat()
    except ValueError: return "INVALID"

def as_int(v, row, label):
    if v is None or s(v) == "": return None
    try: return int(float(v))
    except (TypeError, ValueError):
        err(row, f"{label}은 숫자만 넣어 주세요 (지금 값: {s(v)})")
        return None

def check_privacy(row, text):
    for rx, what in PRIVACY:
        if rx.search(text):
            err(row, f"{what}로 보이는 값이 있습니다. 공개 저장소라 올릴 수 없습니다 → {rx.search(text).group()}")

def slug(gubun, name):
    return hashlib.sha1(f"{gubun}/{name}".encode("utf-8")).hexdigest()[:10]


def read_programs(ws):
    head = [s(c.value) for c in ws[4]]
    need = ["게시","구분","분류","프로그램명","강사","요일","시작시간","종료시간",
            "장소","정원","수강료","재료비","운영시작","운영종료","안내문구"]
    for h in need:
        if h not in head:
            errors.append(f"'프로그램' 시트 4행에 '{h}' 열이 없습니다. 열 이름을 바꾸지 마세요.")
    if errors: return []
    ix = {h: head.index(h) for h in need}
    photo_ix = head.index("사진") if "사진" in head else None   # 사진은 선택 항목

    out, seen = [], {}
    for r in range(5, ws.max_row + 1):
        row = [c.value for c in ws[r]]
        if not any(s(v) for v in row): continue

        pub = s(row[ix["게시"]]).upper()
        if pub in ("X", "×"): continue
        if pub != "O":
            err(r, f"게시 칸은 O 또는 X 만 넣어 주세요 (지금 값: '{s(row[ix['게시']])}')")
            continue

        name  = s(row[ix["프로그램명"]])
        gubun = s(row[ix["구분"]])
        if not name:
            err(r, "프로그램명이 비어 있습니다"); continue
        if gubun not in GUBUN:
            err(r, f"구분은 평생교육/동아리/특강 중 하나여야 합니다 (지금 값: '{gubun}')"); continue

        check_privacy(r, " ".join(s(v) for v in row))

        days = [d for d in re.split(r"[,·/\s]+", s(row[ix["요일"]])) if d]
        if not days:
            err(r, "요일이 비어 있습니다")
        for d in days:
            if d not in DAYS:
                err(r, f"요일이 이상합니다: '{d}' (월·화·수·목·금·토·일 중에서)")

        st, en = as_time(row[ix["시작시간"]]), as_time(row[ix["종료시간"]])
        for label, t in (("시작시간", st), ("종료시간", en)):
            if not TIME_RE.match(t):
                err(r, f"{label}은 09:30 처럼 써 주세요 (지금 값: '{t}')")
        if TIME_RE.match(st) and TIME_RE.match(en) and st >= en:
            err(r, f"종료시간이 시작시간보다 빠르거나 같습니다 ({st} → {en})")

        place = s(row[ix["장소"]])
        if not place: err(r, "장소가 비어 있습니다")

        cap = as_int(row[ix["정원"]], r, "정원")
        fee = as_int(row[ix["수강료"]], r, "수강료") or 0

        d1, d2 = as_date(row[ix["운영시작"]]), as_date(row[ix["운영종료"]])
        for label, d in (("운영시작", d1), ("운영종료", d2)):
            if d == "INVALID": err(r, f"{label} 날짜를 읽을 수 없습니다 (2026-08-19 형식으로)")
        if d1 and d2 and "INVALID" not in (d1, d2) and d1 > d2:
            err(r, f"운영종료가 운영시작보다 빠릅니다 ({d1} → {d2})")

        pid = slug(gubun, name)
        if pid in seen:
            err(r, f"'{name}' 이(가) {seen[pid]}행에도 있습니다. 이름을 다르게 해 주세요 (예: 1반 / 2반)")
        seen[pid] = r

        mat = s(row[ix["재료비"]]) or "없음"
        out.append({
            "id": pid,
            "구분": gubun,
            "분류": s(row[ix["분류"]]) or None,
            "이름": name,
            "강사": s(row[ix["강사"]]) or None,
            "요일": days,
            "시작시간": st,
            "종료시간": en,
            "장소": place,
            "정원": cap,
            "수강료": fee,
            "수강료표시": "무료" if fee == 0 else f"{fee:,}원",
            "재료비별도": (mat == "별도"),
            "운영시작": None if d1 in (None, "INVALID") else d1,
            "운영종료": None if d2 in (None, "INVALID") else d2,
            "안내문구": s(row[ix["안내문구"]]) or None,
            "사진": (s(row[photo_ix]) or None) if photo_ix is not None else None,
        })

    order = {"월":0,"화":1,"수":2,"목":3,"금":4,"토":5,"일":6}
    out.sort(key=lambda p: (order.get(p["요일"][0], 9), p["시작시간"], p["이름"]))
    return out


def read_guide(ws):
    g = {}
    for r in range(5, ws.max_row + 1):
        k = s(ws.cell(r, 1).value)
        if not k: continue
        v = ws.cell(r, 2).value
        g[k] = as_date(v) if isinstance(v, (datetime, date)) else s(v)
    st = g.get("현재상태", "")
    if st not in ("신청기간중", "신청기간아님"):
        errors.append(f"'신청안내' 시트의 현재상태는 '신청기간중' 또는 '신청기간아님' 이어야 합니다 (지금 값: '{st}')")
    check_privacy("신청안내", " ".join(str(v) for v in g.values()))
    return g


def read_notes(ws):
    out = []
    for r in range(5, ws.max_row + 1):
        pub = s(ws.cell(r, 1).value).upper()
        title, body = s(ws.cell(r, 2).value), s(ws.cell(r, 3).value)
        if not title and not body: continue
        if pub in ("X", "×"): continue
        if not title: err(f"안내문 {r}", "제목이 비어 있습니다"); continue
        check_privacy(f"안내문 {r}", title + " " + body)
        out.append({"제목": title, "내용": body})
    return out


def main():
    if not XLSX.exists():
        print(f"::error::{XLSX} 파일이 없습니다.")
        sys.exit(1)

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    for sheet in ("프로그램", "신청안내", "안내문"):
        if sheet not in wb.sheetnames:
            print(f"::error::'{sheet}' 시트가 없습니다. 시트 이름을 바꾸지 마세요.")
            sys.exit(1)

    programs = read_programs(wb["프로그램"])
    guide    = read_guide(wb["신청안내"])
    notes    = read_notes(wb["안내문"])

    if errors:
        print("\n엑셀에서 고쳐야 할 곳이 있습니다. 홈페이지는 그대로 둡니다.\n")
        for e in errors: print("  ✗ " + e)
        print(f"\n모두 {len(errors)}군데입니다. 고쳐서 다시 올려 주세요.")
        print(f"::error::엑셀에 잘못된 값이 {len(errors)}군데 있습니다 (위 목록 참고)")
        sys.exit(1)

    if not programs:
        print("::error::게시할 프로그램이 한 개도 없습니다. 게시 칸이 전부 X 인지 확인해 주세요.")
        sys.exit(1)

    doc = {
        "생성시각": datetime.now(KST).isoformat(timespec="seconds"),
        "출처": "data/북원_홈페이지_프로그램관리.xlsx",
        "학기": guide.get("학기명", ""),
        "문의전화": guide.get("문의전화", "033-747-0516"),
        "신청안내": guide,
        "안내문": notes,
        "프로그램": programs,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding="utf-8")

    by = {}
    for p in programs: by[p["구분"]] = by.get(p["구분"], 0) + 1
    print(f"✓ {OUT} 생성 — 프로그램 {len(programs)}개 (" +
          ", ".join(f"{k} {v}" for k, v in by.items()) + f"), 안내문 {len(notes)}개")
    for w in warns: print("  · " + w)


if __name__ == "__main__":
    main()
