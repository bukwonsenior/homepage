# -*- coding: utf-8 -*-
"""
data/saeop/*.xlsx  ->  data/saeop/<그룹코드>.json   (부서별 파일 → 페이지별 JSON)

부서(사업)마다 엑셀을 따로 둔다. 담당자는 자기 부서 파일만 고친다.
파일 하나에 오류가 있어도 그 페이지만 갱신을 건너뛰고, 나머지는 정상 처리한다.
"""
import json, re, sys
from datetime import datetime, timezone, timedelta
from pathlib import Path
import openpyxl

XDIR = Path("data/saeop")
KST  = timezone(timedelta(hours=9))
SHEET = "사업안내"
DEFAULT_TEL = "033-747-0516"

KINDS = {"부제", "설명", "정보", "절차제목", "절차", "안내", "내용제목", "내용"}
ICON_BY = [("대상","user"),("시간","clock"),("문의","phone"),("전화","phone"),
           ("신청","edit"),("방법","edit"),("접수","edit"),("장소","pin"),("기간","clock")]
PRIVACY = [
    (re.compile(r"\b01[016-9][-. ]?\d{3,4}[-. ]?\d{4}\b"), "휴대전화번호"),
    (re.compile(r"\b\d{6}[-]\d{7}\b"),                     "주민등록번호"),
]

def s(v):
    if v is None: return ""
    return str(v).strip()

def guess_icon(label):
    for key, ic in ICON_BY:
        if key in label: return ic
    return "user"

def parse_file(path):
    """한 엑셀 파일 → (groups dict, errors list)"""
    errors = []
    wb = openpyxl.load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        return {}, [f"'{SHEET}' 시트가 없습니다. 시트 이름을 바꾸지 마세요."]
    ws = wb[SHEET]

    head = [s(c.value) for c in ws[4]]
    need = ["게시", "그룹", "그룹명", "탭", "종류", "라벨/소제목", "내용"]
    for h in need:
        if h not in head:
            errors.append(f"4행에 '{h}' 칸이 없습니다.")
    if errors:
        return {}, errors
    ix = {h: head.index(h) for h in head}

    def col(row, name):
        i = ix.get(name)
        return s(row[i]) if (i is not None and i < len(row)) else ""

    def check_privacy(r, text):
        for rx, what in PRIVACY:
            if rx.search(text):
                errors.append(f"{r}행: {what}로 보이는 값이 있습니다 → {rx.search(text).group()}")

    groups = {}
    for r in range(5, ws.max_row + 1):
        row = [c.value for c in ws[r]]
        if not any(s(v) for v in row): continue
        if col(row, "게시").upper() in ("X", "×"): continue

        grp   = col(row, "그룹")
        gname = col(row, "그룹명")
        tab   = col(row, "탭")
        kind  = col(row, "종류")
        label = col(row, "라벨/소제목")
        text  = col(row, "내용")
        icon  = col(row, "아이콘") if "아이콘" in ix else ""
        tel   = col(row, "전화") if "전화" in ix else ""

        if not grp:  errors.append(f"{r}행: 그룹 칸이 비어 있습니다"); continue
        if not tab:  errors.append(f"{r}행: 탭 칸이 비어 있습니다"); continue
        if kind not in KINDS:
            errors.append(f"{r}행: 종류는 {' / '.join(sorted(KINDS))} 중 하나여야 합니다 (지금: '{kind}')"); continue
        check_privacy(r, f"{label} {text} {tel}")

        g = groups.setdefault(grp, {"대분류": gname or grp, "탭": [], "_ti": {}})
        if gname: g["대분류"] = gname
        if tab not in g["_ti"]:
            g["_ti"][tab] = len(g["탭"])
            g["탭"].append({"이름표": tab, "제목": tab})
        t = g["탭"][g["_ti"][tab]]

        if kind == "부제":      t["부제"] = text
        elif kind == "설명":    t["설명"] = text
        elif kind == "정보":
            if not label: errors.append(f"{r}행: 정보 줄에는 라벨(이용 대상 등)이 필요합니다"); continue
            item = {"라벨": label, "값": text, "icon": icon or guess_icon(label)}
            if tel: item["tel"] = tel
            elif "문의" in label or "전화" in label:
                m = re.search(r"0\d{1,2}[-. )]*\d{3,4}[-. ]*\d{4}", text)
                if m: item["tel"] = m.group()
            t.setdefault("정보", []).append(item)
        elif kind == "절차제목": t["절차제목"] = text
        elif kind == "절차":
            if not label: errors.append(f"{r}행: 절차 줄에는 라벨(단계 제목)이 필요합니다"); continue
            t.setdefault("절차", []).append({"제목": label, "설명": text})
        elif kind == "안내":     t["안내"] = text
        elif kind == "내용제목": t["내용제목"] = text
        elif kind == "내용":
            if not label: errors.append(f"{r}행: 내용 줄에는 라벨(소제목)이 필요합니다"); continue
            arr = t.setdefault("내용", [])
            if arr and arr[-1]["소제목"] == label:
                arr[-1]["항목"].append(text)
            else:
                arr.append({"소제목": label, "항목": [text] if text else []})

    for g in groups.values(): g.pop("_ti", None)
    return groups, errors


def main():
    if not XDIR.exists():
        print(f"::error::{XDIR} 폴더가 없습니다."); sys.exit(1)
    files = sorted(p for p in XDIR.glob("*.xlsx") if not p.name.startswith("~$"))
    if not files:
        print(f"::error::{XDIR} 에 엑셀 파일이 없습니다."); sys.exit(1)

    written = 0
    failed_files = 0
    for path in files:
        try:
            groups, errors = parse_file(path)
        except Exception as e:
            print(f"::error file={path.name}::{path.name} 을(를) 읽지 못했습니다: {e}")
            failed_files += 1
            continue
        if errors:
            failed_files += 1
            print(f"\n[{path.name}] 고쳐야 할 곳 (이 파일은 갱신하지 않음):")
            for e in errors: print("  ✗ " + e)
            print(f"::error file={path.name}::{path.name} 에 잘못된 값 {len(errors)}군데")
            continue
        for code, g in groups.items():
            doc = {
                "생성시각": datetime.now(KST).isoformat(timespec="seconds"),
                "출처": path.name,
                "문의전화": DEFAULT_TEL,
                "대분류": g["대분류"],
                "탭": g["탭"],
            }
            out = XDIR / f"{code}.json"
            out.write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding="utf-8")
            print(f"✓ {out} — 탭 {len(g['탭'])}개")
            written += 1

    print(f"\n완료: {written}개 JSON 생성, 오류 파일 {failed_files}개")
    if written == 0:
        sys.exit(1)


if __name__ == "__main__":
    main()
