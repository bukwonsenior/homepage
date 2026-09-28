# -*- coding: utf-8 -*-
"""
data/연혁.xlsx  ->  data/history.json
한 줄 = 연혁 한 항목. 시트 '연혁', 4행 머리글(게시/연도그룹/연도/날짜/내용), 5행부터 데이터.
연도그룹 = 탭 이름(2026 / 2025 / ... / 1998~2018). 나온 순서대로 탭·항목이 화면에 그려진다.
"""
import json, sys
from datetime import datetime, timezone, timedelta
from pathlib import Path
import openpyxl

SRC  = Path("data/연혁.xlsx")
OUT  = Path("data/history.json")
KST  = timezone(timedelta(hours=9))
SHEET = "연혁"

def s(v):
    return "" if v is None else str(v).strip()

def main():
    if not SRC.exists():
        print(f"::error::{SRC} 파일이 없습니다."); sys.exit(1)
    wb = openpyxl.load_workbook(SRC, data_only=True)
    if SHEET not in wb.sheetnames:
        print(f"::error::'{SHEET}' 시트가 없습니다. 시트 이름을 바꾸지 마세요."); sys.exit(1)
    ws = wb[SHEET]

    head = [s(c.value) for c in ws[4]]
    need = ["게시", "연도그룹", "연도", "날짜", "내용"]
    errors = []
    for h in need:
        if h not in head:
            errors.append(f"4행에 '{h}' 칸이 없습니다.")
    if errors:
        for e in errors: print("::error::" + e)
        sys.exit(1)
    ix = {h: head.index(h) for h in head}
    def col(row, name):
        i = ix.get(name)
        return s(row[i]) if (i is not None and i < len(row)) else ""

    groups = []       # [{이름, 항목:[...]}]
    gmap = {}
    for r in range(5, ws.max_row + 1):
        row = [c.value for c in ws[r]]
        if not any(s(v) for v in row): continue
        if col(row, "게시").upper() in ("X", "×"): continue
        grp  = col(row, "연도그룹")
        year = col(row, "연도")
        date = col(row, "날짜")
        text = col(row, "내용")
        if not grp:  errors.append(f"{r}행: 연도그룹 칸이 비었습니다"); continue
        if not text: errors.append(f"{r}행: 내용 칸이 비었습니다"); continue
        if grp not in gmap:
            gmap[grp] = {"이름": grp, "항목": []}
            groups.append(gmap[grp])
        gmap[grp]["항목"].append({"연도": year, "날짜": date, "내용": text})

    if errors:
        print("\n고쳐야 할 곳:")
        for e in errors: print("  ✗ " + e)
        print(f"::error::잘못된 값 {len(errors)}군데")
        sys.exit(1)

    doc = {
        "생성시각": datetime.now(KST).isoformat(timespec="seconds"),
        "출처": SRC.name,
        "그룹": groups,
    }
    OUT.write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding="utf-8")
    total = sum(len(g["항목"]) for g in groups)
    print(f"✓ {OUT} — 탭 {len(groups)}개, 항목 {total}개")

if __name__ == "__main__":
    main()
