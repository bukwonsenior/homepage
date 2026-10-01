# -*- coding: utf-8 -*-
"""
시설안내 엑셀(북원_시설안내.xlsx) → facility.json

'시설안내' 시트, 4행 머리글: 게시 | 층명 | 실명 | 사진 | 역할 | 이용방법   (5행부터 데이터)
- 층명이 같은 줄끼리 한 층으로 묶이고, 나온 순서대로 화면에 그려진다.
  층명 '1층'→ 배지 1F, '지하1층'/'B1층' → B1, 그 밖(예: 전경)은 그대로 배지.
- 사진: img/facility/ 폴더의 파일명. 쉼표로 여러 장(최대 3). 확장자를 빼먹으면 .jpg로 본다.
- 선택 시트 '기관정보'(항목|내용) → 시설안내 맨 위 기관 정보.
- (예전 형식의 '층코드' 칸이 있으면 그 값을 층 구분에 그대로 쓴다)
"""
import json, re, sys
from pathlib import Path
import openpyxl

XLSX = Path("북원_시설안내.xlsx")
OUT  = Path("facility.json")
IMG_BASE = "img/facility/"
IMG_DIR  = Path("img/facility")

errors, warns = [], []
def s(v): return "" if v is None else str(v).strip()
JUMIN = re.compile(r"\b\d{6}-\d{7}\b")

def floor_code(name):
    m = re.fullmatch(r"(\d+)\s*층", name)
    if m: return "f" + m.group(1), m.group(1) + "F"
    m = re.fullmatch(r"(?:지하|B)\s*(\d+)\s*층?", name, re.I)
    if m: return "b" + m.group(1), "B" + m.group(1)
    if name == "전경": return "exterior", "전경"
    return name, name

def main():
    if not XLSX.exists(): print(f"::error::{XLSX} 파일이 없습니다."); sys.exit(1)
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    if "시설안내" not in wb.sheetnames: print("::error::'시설안내' 시트가 없습니다."); sys.exit(1)
    ws = wb["시설안내"]
    head = [s(c.value) for c in ws[4]]
    for h in ["게시", "층명", "실명", "사진", "역할", "이용방법"]:
        if h not in head: errors.append(f"4행에 '{h}' 칸이 없습니다.")
    if errors:
        for e in errors: print("::error::" + e)
        sys.exit(1)
    ix = {h: i for i, h in enumerate(head) if h}
    def col(row, n):
        i = ix.get(n); return s(row[i]) if (i is not None and i < len(row)) else ""

    floors, order = {}, []
    for r in range(5, ws.max_row + 1):
        row = [c.value for c in ws[r]]
        if not any(s(v) for v in row): continue
        pub = col(row, "게시").upper()
        if pub in ("X", "×"): continue
        fname, rname = col(row, "층명"), col(row, "실명")
        if not fname: errors.append(f"{r}행: 층명이 비어 있습니다"); continue
        if not rname: errors.append(f"{r}행: 실명이 비어 있습니다"); continue
        if JUMIN.search(" ".join(s(v) for v in row)): errors.append(f"{r}행: 주민등록번호로 보이는 값이 있습니다")
        code, badge = floor_code(fname)
        if "층코드" in ix and col(row, "층코드"): code = col(row, "층코드")
        photos = []
        for p in [p.strip() for p in col(row, "사진").split(",") if p.strip()][:3]:
            if p.startswith("http"): photos.append(p); continue
            if not re.search(r"\.(jpe?g|png|webp|gif)$", p, re.I): p += ".jpg"
            if IMG_DIR.exists() and not (IMG_DIR / p).exists():
                warns.append(f"{r}행 {rname}: img/facility/{p} 사진 파일이 아직 없습니다 (화면엔 '사진 준비 중')")
            photos.append(IMG_BASE + p)
        if code not in floors:
            floors[code] = {"code": code, "badge": badge, "name": fname, "rooms": []}
            order.append(code)
        floors[code]["rooms"].append({"name": rname, "photos": photos, "role": col(row, "역할"), "usage": col(row, "이용방법")})

    if errors:
        print("\n엑셀에서 고쳐야 할 곳:")
        for e in errors: print("  ✗ " + e)
        print(f"::error::엑셀 오류 {len(errors)}건"); sys.exit(1)
    for w in warns: print("::warning::" + w)
    info = []
    if "기관정보" in wb.sheetnames:   # 선택 시트: 맨 위 기관 정보 (항목 | 내용)
        wi = wb["기관정보"]
        for r in range(5, wi.max_row + 1):
            k, v = s(wi.cell(r, 1).value), s(wi.cell(r, 2).value)
            if k or v: info.append({"항목": k, "내용": v})
    result = [floors[c] for c in order]
    doc = {"기관정보": info, "층": result}
    OUT.write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"✓ {OUT} 생성 — {len(result)}개 층, {sum(len(f['rooms']) for f in result)}개 실")

if __name__ == "__main__":
    main()
