# -*- coding: utf-8 -*-
"""
data/자원봉사.xlsx -> data/volunteer.json
시트 '자원봉사', 4행 머리글(게시/종류/이름·라벨/내용/태그/아이콘), 5행부터 데이터.
종류: 소개 / 신청 / 절차 / 일반봉사 / 시니어봉사 / 혜택 / 안내
"""
import json, sys
from datetime import datetime, timezone, timedelta
from pathlib import Path
import openpyxl

SRC = Path("data/자원봉사.xlsx")
OUT = Path("data/volunteer.json")
KST = timezone(timedelta(hours=9))
SHEET = "자원봉사"
KINDS = {"소개","신청","절차","일반봉사","시니어봉사","혜택","안내"}

def s(v): return "" if v is None else str(v).strip()

def main():
    if not SRC.exists():
        print(f"::error::{SRC} 파일이 없습니다."); sys.exit(1)
    wb = openpyxl.load_workbook(SRC, data_only=True)
    if SHEET not in wb.sheetnames:
        print(f"::error::'{SHEET}' 시트가 없습니다."); sys.exit(1)
    ws = wb[SHEET]
    head = [s(c.value) for c in ws[4]]
    need = ["게시","종류","이름/라벨","내용","태그","아이콘"]
    errors=[]
    for h in need:
        if h not in head: errors.append(f"4행에 '{h}' 칸이 없습니다.")
    if errors:
        for e in errors: print("::error::"+e)
        sys.exit(1)
    ix={h:head.index(h) for h in head}
    def col(row,name):
        i=ix.get(name); return s(row[i]) if (i is not None and i<len(row)) else ""

    doc={"생성시각":datetime.now(KST).isoformat(timespec="seconds"),"출처":SRC.name,
         "소개":"","신청":[],"절차":[],"일반봉사":[],"시니어봉사":[],"혜택":[],"안내":""}
    for r in range(5, ws.max_row+1):
        row=[c.value for c in ws[r]]
        if not any(s(v) for v in row): continue
        if col(row,"게시").upper() in ("X","×"): continue
        kind=col(row,"종류"); name=col(row,"이름/라벨"); text=col(row,"내용")
        tag=col(row,"태그"); icon=col(row,"아이콘")
        if kind not in KINDS:
            errors.append(f"{r}행: 종류는 {' / '.join(sorted(KINDS))} 중 하나 (지금:'{kind}')"); continue
        if kind=="소개": doc["소개"]=text
        elif kind=="안내": doc["안내"]=text
        elif kind=="신청": doc["신청"].append({"라벨":name,"값":text,"아이콘":icon})
        elif kind=="절차": doc["절차"].append(name or text)
        elif kind=="혜택": doc["혜택"].append(text or name)
        elif kind=="일반봉사": doc["일반봉사"].append({"이름":name,"설명":text,"태그":tag,"아이콘":icon})
        elif kind=="시니어봉사": doc["시니어봉사"].append({"이름":name,"설명":text,"아이콘":icon})
    if errors:
        for e in errors: print("  ✗ "+e)
        print(f"::error::잘못된 값 {len(errors)}군데"); sys.exit(1)
    OUT.write_text(json.dumps(doc,ensure_ascii=False,indent=1),encoding="utf-8")
    print(f"✓ {OUT} — 일반 {len(doc['일반봉사'])} / 시니어 {len(doc['시니어봉사'])} / 절차 {len(doc['절차'])} / 혜택 {len(doc['혜택'])}")

if __name__=="__main__":
    main()
